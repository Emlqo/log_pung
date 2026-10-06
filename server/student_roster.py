"""Bounded XLSX import and conservative name matching. Never retain raw workbooks."""
from collections import Counter
from io import BytesIO
import re
import unicodedata
from zipfile import ZipFile, BadZipFile

MAX_ROSTER_BYTES = 2 * 1024 * 1024

def name_key(value):
    return re.sub(r'\s+', '', unicodedata.normalize('NFC', value or ''))

def parse_roster(content, year):
    if not content or len(content) > MAX_ROSTER_BYTES:
        raise ValueError('명렬표는 2MB 이하 XLSX 파일로 올려주세요.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 200 or sum(x.file_size for x in entries) > 20 * 1024 * 1024:
                raise ValueError('압축을 푼 명렬표가 너무 큽니다.')
            if any(x.filename.endswith('vbaProject.bin') for x in entries):
                raise ValueError('매크로가 없는 XLSX 파일이 필요합니다.')
    except (BadZipFile, OSError):
        raise ValueError('정상적인 XLSX 파일이 아닙니다.') from None
    from openpyxl import load_workbook
    book = None
    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        rows = []; seats = set(); years = set(); scanned = 0
        for sheet in book:
            if sheet.max_row and sheet.max_row > 10000 or sheet.max_column and sheet.max_column > 100:
                raise ValueError('명렬표는 시트당 10,000행·100열 이내여야 합니다.')
            columns = None
            for values in sheet.iter_rows(max_col=100, values_only=True):
                scanned += 1
                if scanned > 10000: raise ValueError('명렬표는 총 10,000행 이내여야 합니다.')
                cells = [str(x).strip() if x is not None else '' for x in values]
                for cell in cells:
                    match = re.search(r'(20\d{2})\s*학년도', cell)
                    if match: years.add(int(match[1]))
                if all(label in cells for label in ('학년','반','번호','성명')):
                    columns = [cells.index(label) for label in ('학년','반','번호','성명')]
                    continue
                if columns is None: continue
                grade, classroom, number, name = [cells[i] for i in columns]
                if not any((grade, classroom, number, name)): continue
                # Footer / title rows have neither a student number nor a grade.
                if not number and not re.fullmatch(r'\d+\s*(?:학년)?', grade): continue
                def integer(value, suffix, maximum):
                    match = re.fullmatch(r'(\d+)\s*'+suffix+'?', value)
                    if not match or not 1 <= int(match[1]) <= maximum:
                        raise ValueError(f'{scanned}행의 학년·반·번호를 확인해주세요.')
                    return int(match[1])
                g = integer(grade, '(?:학년)', 12)
                c = integer(classroom, '(?:반)', 99)
                n = integer(number, '(?:번)', 999)
                if not name or len(name) > 200 or any(ord(x)<32 for x in name):
                    raise ValueError(f'{scanned}행의 성명을 확인해주세요.')
                seat = (g,c,n)
                if seat in seats: raise ValueError('학년·반·번호가 중복된 행이 있습니다.')
                seats.add(seat)
                rows.append(dict(grade=g,classroom=c,number=n,name=name))
        if years and years != {year}: raise ValueError('선택한 학년도와 파일의 학년도가 다릅니다.')
        if not rows: raise ValueError('학년·반·번호·성명 열과 학생 행을 찾지 못했습니다.')
        return rows
    except ValueError: raise
    except Exception: raise ValueError('명렬표를 읽을 수 없습니다. XLSX 양식을 확인해주세요.') from None
    finally:
        if book: book.close()

def match_students(known, directory, aliases, roster, overrides):
    """Only observed students are returned; ambiguity includes unseen directory users."""
    directory_counts = Counter(name_key(name) for name in directory.values())
    roster_counts = Counter(name_key(row['name']) for row in roster)
    by_name = {name_key(row['name']): row for row in roster}
    occupied = {(r['grade'],r['classroom'],r['number']): email for email,r in overrides.items() if r['number'] is not None}
    result = []
    for person in known:
        email = person['email']; name = directory.get(email,''); key = name_key(name)
        item = {**person,'display_name':name,'alias':aliases.get(email,''),
                'grade':None,'classroom':None,'number':None,'source':'unassigned','reason':''}
        if email in overrides:
            item.update(overrides[email],source='manual')
        elif not key: item['reason']='이메일에 연결된 이름이 없습니다.'
        elif directory_counts[key] > 1 or roster_counts[key] > 1: item['reason']='동명이인 · 직접 반을 지정해주세요.'
        elif key not in by_name: item['reason']='명렬표에서 이름을 찾지 못했습니다.'
        else:
            row = by_name[key]; seat = (row['grade'],row['classroom'],row['number'])
            if seat in occupied and occupied[seat] != email:
                item['reason']='같은 번호에 수동 배정된 학생이 있습니다.'
            else: item.update({k:row[k] for k in ('grade','classroom','number')},source='auto')
        result.append(item)
    return result
